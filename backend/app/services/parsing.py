"""File parsing for PDF, DOCX, and TXT uploads.

This module is intentionally narrow: it accepts a path on disk plus a content
type / extension hint, and returns plain text. Errors are wrapped in
``ParsingError`` to give the API a stable failure surface.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


class ParsingError(RuntimeError):
    """Raised when a document cannot be parsed."""


def extract_text(path: Path, extension: str) -> str:
    """Extract plain text from a supported file type.

    ``extension`` is the lowercase extension without the dot (``"pdf"``,
    ``"docx"``, or ``"txt"``).
    """
    ext = extension.lower().lstrip(".")
    try:
        if ext == "pdf":
            return _extract_pdf(path)
        if ext == "docx":
            return _extract_docx(path)
        if ext == "txt":
            return _extract_txt(path)
    except ParsingError:
        raise
    except Exception as exc:  # pragma: no cover - defensive
        raise ParsingError(f"Failed to parse .{ext} file: {exc}") from exc
    raise ParsingError(f"Unsupported file extension: .{ext}")


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            # Some pages legitimately fail to extract (images, encrypted form
            # fields, etc.). Skip them rather than aborting the whole parse.
            pages.append("")
    text = "\n\n".join(p.strip() for p in pages if p and p.strip())
    if not text.strip():
        raise ParsingError("PDF contained no extractable text")
    return text


def _extract_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    parts.append(text)
    text = "\n".join(parts)
    if not text.strip():
        raise ParsingError("DOCX contained no extractable text")
    return text


def _extract_txt(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 should always succeed
        raise ParsingError("Could not decode TXT file with any known encoding")
    if not text.strip():
        raise ParsingError("TXT file was empty")
    return text
